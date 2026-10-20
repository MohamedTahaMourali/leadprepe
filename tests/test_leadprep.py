import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from leadprep import agents, cleaning as c, db, queries as q
from leadprep.excel_export import write_workbook

COLS = c.EXPECTED


def row(**kw):
    base = {k: "" for k in COLS}
    base.update(siren="123456789", nom="X", secteur="compta", code_postal="69003", effectif="NN",
                entrepreneur_individuel="non")
    base.update(kw)
    return base


class CleaningTests(unittest.TestCase):
    def build(self, *rows, **kw):
        df, rej = c.clean(pd.DataFrame(rows, columns=COLS), "t.csv", **kw)
        return df

    def test_segments_partition(self):
        rows = [
            row(siren="000000001", site_web="https://a.fr", emails="x@a.fr", emails_valides="x@a.fr", telephones="04 78 12 34 56"),
            row(siren="000000002", site_web="https://b.fr", telephones="04 78 12 34 56"),
            row(siren="000000003", site_web="https://c.fr", emails="y@c.fr"),
            row(siren="000000004", site_web="https://d.fr"),
            row(siren="000000005"),
            row(siren="000000006", emails="z@e.fr", telephones="04 78 12 34 56"),   # tel+email sans site
        ]
        got = dict(zip(self.build(*rows).siren, self.build(*rows).segment))
        self.assertEqual(got, {"000000001": "complet", "000000002": "tel_sans_email", "000000003": "email_sans_tel",
                               "000000004": "site_seul", "000000005": "autres", "000000006": "tel_email_sans_site"})

    def test_strict_email_uses_valid_mx(self):
        r = row(site_web="https://a.fr", emails="x@a.fr", emails_valides="", telephones="04 78 12 34 56")
        self.assertEqual(self.build(r).segment[0], "complet")
        self.assertEqual(self.build(r, strict_email=True).segment[0], "tel_sans_email")

    def test_siren_padding_and_rejects(self):
        df, rej = c.clean(pd.DataFrame([row(siren="12345678"), row(siren="abc")], columns=COLS), "t.csv")
        self.assertEqual(df.siren[0], "012345678")
        self.assertEqual(rej, 1)

    def test_email_choice_prefers_valid_own_domain_over_third_party(self):
        r = row(site_web="https://www.cabinet.fr", emails="x@oracle.com | contact@cabinet.fr | y@gmail.com",
                emails_valides="x@oracle.com | contact@cabinet.fr | y@gmail.com", telephones="")
        d = self.build(r).iloc[0]
        self.assertEqual(d.email_principal, "contact@cabinet.fr")
        self.assertEqual((d.email_type, d.email_domaine_proprio, d.email_tiers), ("generique", 1, 0))

    def test_phones_and_suspect(self):
        many = " | ".join(f"04 78 12 34 {i:02d}" for i in range(10, 17))
        d = self.build(row(telephones="01 23 45 67 89 | 06 12 34 56 78"), row(siren="000000002", telephones=many))
        self.assertEqual(d.telephone_e164[0], "+33123456789")
        self.assertEqual(d.telephone_mobile[0], "06 12 34 56 78")
        self.assertEqual(int(d.telephone_suspect[1]), 1)

    def test_departement(self):
        self.assertEqual([c.departement(x) for x in ("69003", "97100", "20090", "20600", "")], ["69", "971", "2A", "2B", ""])

    def test_score_bounds(self):
        d = self.build(row(site_web="https://a.fr", site_confiance="SIREN", emails="x@a.fr", emails_valides="x@a.fr",
                           telephones="04 78 12 34 56", effectif="20 à 49"))
        self.assertEqual(d.lead_score[0], 100)
        self.assertEqual(self.build(row()).lead_score[0], 10)    # non-EI sans contact


class DbTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.conn = db.connect(self.tmp / "t.db")
        db.init_db(self.conn)
        self.csv = self.tmp / "01.csv"
        pd.DataFrame([row(siren="000000001", site_web="https://a.fr", telephones="04 78 12 34 56"),
                      row(siren="000000002", nom="=cmd()")], columns=COLS).to_csv(self.csv, sep=";", index=False)

    def test_idempotent_and_upsert(self):
        self.assertFalse(db.ingest_file(self.conn, self.csv)["skipped"])
        self.assertTrue(db.ingest_file(self.conn, self.csv)["skipped"])
        r = db.ingest_file(self.conn, self.csv, force=True)
        self.assertEqual((r["nouveaux"], r["mis_a_jour"]), (0, 2))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM leads").fetchone()[0], 2)

    def test_poorer_row_does_not_overwrite_richer(self):
        db.ingest_file(self.conn, self.csv)
        poor = self.tmp / "02.csv"
        pd.DataFrame([row(siren="000000001")], columns=COLS).to_csv(poor, sep=";", index=False)
        db.ingest_file(self.conn, poor)
        self.assertEqual(self.conn.execute("SELECT segment FROM leads WHERE siren='000000001'").fetchone()[0], "tel_sans_email")

    def test_merge_file_is_ignored(self):
        (self.tmp / "_france_all.csv").write_text("x")
        self.assertEqual([p.name for p in db.discover_csv([self.tmp])], ["01.csv"])

    def test_queue_and_results(self):
        db.ingest_file(self.conn, self.csv)
        self.assertEqual(list(agents.queue(self.conn, "phone_caller").siren), ["000000001"])
        agents.record_result(self.conn, "000000001", "phone_caller", "done", {"ok": 1})
        self.assertTrue(agents.queue(self.conn, "phone_caller").empty)

    def test_excel_formula_safe_and_sheets(self):
        db.ingest_file(self.conn, self.csv)
        df = q.leads_page(self.conn)
        out = self.tmp / "o.xlsx"
        write_workbook(out, df, "t")
        import openpyxl
        wb = openpyxl.load_workbook(out)
        self.assertIn("5_Autres", wb.sheetnames)
        cell = next(r for r in wb["5_Autres"].iter_rows(min_row=2, values_only=True) if r[0] == "000000002")
        self.assertEqual(cell[1], "=cmd()")
        self.assertNotEqual(wb["5_Autres"]["B2"].data_type, "f")

    def test_sector_priority_runs(self):
        db.ingest_file(self.conn, self.csv)
        self.assertEqual(len(q.sector_priority(self.conn)), 1)


if __name__ == "__main__":
    unittest.main()
